from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.comments import Comment
FONT='Arial'
BLUE=Font(name=FONT,color='0000FF',size=10)
BLACK=Font(name=FONT,color='000000',size=10)
GREEN=Font(name=FONT,color='008000',size=10)
BOLD=Font(name=FONT,bold=True,size=10)
H1=Font(name=FONT,bold=True,size=14)
H2=Font(name=FONT,bold=True,size=11,color='FFFFFF')
NOTE=Font(name=FONT,italic=True,size=9,color='555555')
YELLOW=PatternFill('solid',fgColor='FFFF00')
HEAD=PatternFill('solid',fgColor='1F3A68')
LIGHT=PatternFill('solid',fgColor='EEF2F8')
THIN=Border(bottom=Side(style='thin',color='BBBBBB'))
USD='$#,##0;($#,##0);-'
USDK='$#,##0,"k";($#,##0,"k");-'
USDM='$#,##0.0,,"M";($#,##0.0,,"M");-'
NUM='#,##0;(#,##0);-'
NUM1='#,##0.0;(#,##0.0);-'
PCT='0.0%;(0.0%);-'
MULT='0.00"x"'
def w(ws,ref,val,font=BLACK,fmt=None,fill=None,bold=False,comment=None,wrap=False,align=None):
    c=ws[ref]; c.value=val; c.font=Font(name=FONT,bold=bold or font.bold,italic=font.italic,size=font.size,color=font.color)
    if fmt: c.number_format=fmt
    if fill: c.fill=fill
    if comment: c.comment=Comment(comment,'Agentic Office model')
    if wrap or align: c.alignment=Alignment(wrap_text=wrap,horizontal=align,vertical='top')
    return c
def header(ws,row,labels,start=1):
    for i,l in enumerate(labels):
        c=ws.cell(row=row,column=start+i,value=l); c.font=H2; c.fill=HEAD; c.alignment=Alignment(wrap_text=True,vertical='center')
def widths(ws,d):
    for k,v in d.items(): ws.column_dimensions[k].width=v

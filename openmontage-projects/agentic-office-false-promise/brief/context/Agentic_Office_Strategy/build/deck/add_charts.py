import json, sys
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION, XL_TICK_LABEL_POSITION, XL_MARKER_STYLE
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
src, out = sys.argv[1], sys.argv[2]
prs = Presentation(src); specs = json.load(open('charts.json'))
H = 'Helvetica Neue'
def rgb(h): return RGBColor.from_string(h)
for sp in specs:
    slide = prs.slides[sp['slide'] - 1]; o = sp['opts']; data = sp['data']
    cd = CategoryChartData(); cd.categories = data[0]['labels']
    for ser in data: cd.add_series(ser['name'], ser['values'])
    if sp['type'] == 'line': ct = XL_CHART_TYPE.LINE_MARKERS
    else: ct = XL_CHART_TYPE.BAR_CLUSTERED if o.get('barDir') == 'bar' else XL_CHART_TYPE.COLUMN_CLUSTERED
    gf = slide.shapes.add_chart(ct, Inches(o['x']), Inches(o['y']), Inches(o['w']), Inches(o['h']), cd); ch = gf.chart
    ch.font.name = H; ch.font.size = Pt(o.get('catAxisLabelFontSize', 10)); ch.font.color.rgb = rgb('1B1B1A')
    if o.get('showTitle'):
        ch.has_title = True; tf = ch.chart_title.text_frame; tf.text = o['title']
        r = tf.paragraphs[0].runs[0]; r.font.size = Pt(o.get('titleFontSize', 10)); r.font.name = H; r.font.color.rgb = rgb(o.get('titleColor', '7A776F')); r.font.bold = False
    else: ch.has_title = False
    ch.has_legend = bool(o.get('showLegend'))
    if ch.has_legend:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False; ch.legend.font.size = Pt(o.get('legendFontSize', 10)); ch.legend.font.name = H
    plot = ch.plots[0]
    if sp['type'] != 'line':
        plot.gap_width = o.get('barGapWidthPct', 60)
        cols = o.get('chartColors', ['3157F6'])
        ser = plot.series[0]
        for i in range(len(data[0]['values'])):
            pt = ser.points[i]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = rgb(cols[i] if len(cols) > 1 and i < len(cols) else cols[0]); pt.format.line.fill.background()
    else:
        cols = o.get('chartColors')
        for i, ser in enumerate(plot.series):
            ser.format.line.color.rgb = rgb(cols[i]); ser.format.line.width = Pt(2.25); ser.smooth = False
            ser.marker.style = XL_MARKER_STYLE.CIRCLE; ser.marker.size = 6; ser.marker.format.fill.solid(); ser.marker.format.fill.fore_color.rgb = rgb(cols[i]); ser.marker.format.line.color.rgb = rgb(cols[i])
    if o.get('showValue'):
        plot.has_data_labels = True; dl = plot.data_labels
        dl.font.size = Pt(o.get('dataLabelFontSize', 10)); dl.font.name = H; dl.font.color.rgb = rgb('1B1B1A')
        if o.get('dataLabelFormatCode'): dl.number_format = o['dataLabelFormatCode']; dl.number_format_is_linked = False
        dl.position = XL_LABEL_POSITION.OUTSIDE_END
    va = ch.value_axis; ca = ch.category_axis
    va.has_major_gridlines = 'valGridLine' in o and o['valGridLine'].get('style') != 'none'
    if va.has_major_gridlines: va.major_gridlines.format.line.color.rgb = rgb('E6E3DC'); va.major_gridlines.format.line.width = Pt(0.5)
    ca.has_major_gridlines = False
    if 'valAxisMinVal' in o: va.minimum_scale = o['valAxisMinVal']
    if 'valAxisMaxVal' in o: va.maximum_scale = o['valAxisMaxVal']
    if o.get('valAxisHidden'):
        from pptx.enum.chart import XL_TICK_MARK
        va.tick_label_position = XL_TICK_LABEL_POSITION.NONE; va.major_tick_mark = XL_TICK_MARK.NONE; va.minor_tick_mark = XL_TICK_MARK.NONE
        va.format.line.fill.background()
    else:
        va.tick_labels.font.size = Pt(o.get('valAxisLabelFontSize', 9)); va.tick_labels.font.color.rgb = rgb('7A776F')
        if o.get('valAxisLabelFormatCode'): va.tick_labels.number_format = o['valAxisLabelFormatCode']; va.tick_labels.number_format_is_linked = False
        va.format.line.fill.background()
    ca.tick_labels.font.size = Pt(o.get('catAxisLabelFontSize', 10)); ca.tick_labels.font.name = H
    ca.format.line.color.rgb = rgb('D9D6CE')
    ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    if o.get('catAxisOrientation') == 'maxMin': ca.reverse_order = True
prs.save(out); print('charts added:', len(specs))
